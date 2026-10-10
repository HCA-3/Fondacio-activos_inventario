from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.tiempo import hoy
from tests.ayudas import nueva_sede, nuevo_equipo

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def admin(cliente_admin: TestClient) -> TestClient:
    nueva_sede(cliente_admin)
    return cliente_admin


def descargar(cliente: TestClient, **filtros) -> dict[str, list[dict]]:
    """Baja el Excel y lo devuelve como {hoja: [fila como diccionario]}."""
    respuesta = cliente.get("/api/equipos/export.xlsx", params=filtros)
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.headers["content-type"] == XLSX
    libro = load_workbook(BytesIO(respuesta.content))
    hojas = {}
    for hoja in libro.worksheets:
        encabezados, *filas = hoja.iter_rows(values_only=True)
        hojas[hoja.title] = [dict(zip(encabezados, fila, strict=True)) for fila in filas]
    return hojas


def test_el_excel_trae_las_mismas_hojas_y_columnas_que_generaba_el_frontend(
    admin: TestClient,
) -> None:
    nuevo_equipo(
        admin,
        "PC-01",
        computerType="PORTÁTIL",
        assignedTo="Coordinación",
        networkHostname="FND-PC01",
        os="Windows 11 Pro",
        processor="Intel Core i5-7300U",
        ram="16 GB",
        disk1Capacity="256 GB",
        disk1Tech="SSD NVMe",
        ip="192.168.1.20",
        recommendations="Cambiar batería",
    )
    nuevo_equipo(admin, "TAB-01", category="tablet", model="Tablet Educativa", os="Android")

    hojas = descargar(admin)

    assert list(hojas) == [
        "Censo General de Equipos",
        "Detalle Hardware PCs",
        "Software y Mantenimiento",
        "Dispositivos Móviles (Tablets)",
    ]
    censo, tablet = hojas["Censo General de Equipos"]
    assert censo == {
        "ID Equipo": "PC-01",
        "Tipo Dispositivo": "PORTÁTIL",
        "Marca / Fabricante": "Dell",
        "Modelo": "Latitude 7480",
        "Nombre del Equipo": "FND-PC01",
        "Cuenta / Usuario": "Coordinación",
        "Sistema Operativo": "Windows 11 Pro",
        "Versión / Edición": "Windows 11 Pro",
        "Estado Licencia": "Activa",
        "Procesador (CPU)": "Intel Core i5-7300U",
        "RAM Instalada": "16 GB",
        "Almacenamiento (Disco)": "256 GB SSD NVMe",
        "Tarjeta Gráfica (GPU)": "GPU Integrada",
        "Dirección IP (Red)": "192.168.1.20 (ALTINET FONDACIO)",
        "Estado Operativo / Mantenimiento": "Cambiar batería",
    }
    assert (tablet["ID Equipo"], tablet["Tipo Dispositivo"]) == ("TAB-01", "Tablet")
    # Los computadores y las tablets van en hojas distintas.
    assert [f["ID Equipo"] for f in hojas["Detalle Hardware PCs"]] == ["PC-01"]
    assert len(hojas["Detalle Hardware PCs"][0]) == 32
    assert hojas["Software y Mantenimiento"][0]["Versión / Build"] == "22H2"
    assert [f["ID Dispositivo"] for f in hojas["Dispositivos Móviles (Tablets)"]] == ["TAB-01"]


def test_el_excel_respeta_los_filtros_de_la_lista(admin: TestClient) -> None:
    nueva_sede(admin, "YLDC Potosí")
    nuevo_equipo(admin, "PC-02", status="bodega")
    nuevo_equipo(admin, "PC-01")
    nuevo_equipo(admin, "PC-03", sede="YLDC Potosí", serial="XYZ")

    def placas(**filtros) -> list[str]:
        return [f["ID Equipo"] for f in descargar(admin, **filtros)["Censo General de Equipos"]]

    assert placas() == ["PC-01", "PC-02", "PC-03"]
    assert placas(sede="YLDC Potosí") == ["PC-03"]
    assert placas(estado="bodega") == ["PC-02"]
    assert placas(q="xyz") == ["PC-03"]


def test_sin_equipos_el_excel_solo_trae_los_encabezados_del_censo(admin: TestClient) -> None:
    assert descargar(admin) == {"Censo General de Equipos": []}


def test_un_texto_que_empieza_por_igual_no_se_vuelve_formula(admin: TestClient) -> None:
    nuevo_equipo(admin, "PC-01", brand='=HYPERLINK("http://malo","clic")')

    contenido = admin.get("/api/equipos/export.xlsx").content
    celda = load_workbook(BytesIO(contenido))["Censo General de Equipos"]["C2"]

    assert celda.value == '=HYPERLINK("http://malo","clic")'
    assert celda.data_type == "s"


def test_el_archivo_se_descarga_con_nombre_y_fecha(admin: TestClient) -> None:
    cabecera = admin.get("/api/equipos/export.xlsx").headers["content-disposition"]

    assert (
        cabecera == f'attachment; filename="Inventario_Equipos_Fondacio_{hoy().isoformat()}.xlsx"'
    )


def test_el_rol_consulta_descarga_el_excel(admin: TestClient, cliente_consulta: TestClient) -> None:
    assert cliente_consulta.get("/api/equipos/export.xlsx").status_code == 200


def test_sin_sesion_no_se_descarga_el_excel(cliente: TestClient) -> None:
    assert cliente.get("/api/equipos/export.xlsx").status_code == 401
